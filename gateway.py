from mitmproxy import http

import json
import uuid
import requests

from datetime import datetime, timezone
from io import BytesIO
from email.parser import BytesParser
from email.policy import default

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook


# ============================================================
# Gateway 설정
# ============================================================

OPENAI_HOST = "api.openai.com"

OPENAI_RESPONSES_ENDPOINT = "/v1/responses"
OPENAI_FILES_ENDPOINT = "/v1/files"

# 다음 모듈
NEXT_MODULE_URL = (
    "http://127.0.0.1:3000/internal/gateway/events"
)


# ============================================================
# Client IP
# ============================================================

def get_client_ip(flow: http.HTTPFlow) -> str | None:

    if flow.client_conn.peername:
        return flow.client_conn.peername[0]

    return None


# ============================================================
# 파일 텍스트 추출
# TXT / PDF / DOCX / XLSX
# ============================================================

def extract_file_text(
    filename: str | None,
    file_bytes: bytes,
) -> str | None:

    if not filename:
        return None

    extension = filename.lower().rsplit(".", 1)[-1]

    # --------------------------------------------------------
    # TXT
    # --------------------------------------------------------

    if extension == "txt":

        try:
            return file_bytes.decode(
                "utf-8",
                errors="ignore",
            )

        except Exception as e:

            print(
                "❌ TXT text extraction failed:",
                e,
            )

            return None

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if extension == "pdf":

        try:

            reader = PdfReader(
                BytesIO(file_bytes)
            )

            texts = []

            for page in reader.pages:

                text = page.extract_text()

                if text:
                    texts.append(text)

            return "\n".join(texts)

        except Exception as e:

            print(
                "❌ PDF text extraction failed:",
                e,
            )

            return None

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    if extension == "docx":

        try:

            document = Document(
                BytesIO(file_bytes)
            )

            texts = []

            # 일반 문단
            for paragraph in document.paragraphs:

                if paragraph.text:
                    texts.append(
                        paragraph.text
                    )

            # 표 내부 텍스트
            for table in document.tables:

                for row in table.rows:

                    values = []

                    for cell in row.cells:

                        if cell.text:
                            values.append(
                                cell.text
                            )

                    if values:
                        texts.append(
                            " | ".join(values)
                        )

            return "\n".join(texts)

        except Exception as e:

            print(
                "❌ DOCX text extraction failed:",
                e,
            )

            return None

    # --------------------------------------------------------
    # XLSX
    # --------------------------------------------------------

    if extension == "xlsx":

        try:

            workbook = load_workbook(
                BytesIO(file_bytes),
                read_only=True,
                data_only=True,
            )

            texts = []

            for sheet in workbook.worksheets:

                # Sheet 이름도 정보로 포함
                texts.append(
                    f"[Sheet: {sheet.title}]"
                )

                for row in sheet.iter_rows(
                    values_only=True
                ):

                    values = [
                        str(value)
                        for value in row
                        if value is not None
                    ]

                    if values:

                        texts.append(
                            " | ".join(values)
                        )

            return "\n".join(texts)

        except Exception as e:

            print(
                "❌ XLSX text extraction failed:",
                e,
            )

            return None

    # --------------------------------------------------------
    # 지원하지 않는 형식
    # --------------------------------------------------------

    print(
        f"⚠️ Unsupported file type: {filename}"
    )

    return None


# ============================================================
# Responses API 콘텐츠 추출
# ============================================================

def extract_content(
    body: dict,
) -> list[dict]:

    results = []

    input_data = body.get("input")

    # --------------------------------------------------------
    # input = 문자열
    # --------------------------------------------------------

    if isinstance(input_data, str):

        results.append({
            "type": "prompt",
            "content": input_data,
        })

    # --------------------------------------------------------
    # input = 배열
    # --------------------------------------------------------

    elif isinstance(input_data, list):

        for item in input_data:

            # 문자열
            if isinstance(item, str):

                results.append({
                    "type": "prompt",
                    "content": item,
                })

                continue

            # 객체가 아니면 무시
            if not isinstance(item, dict):
                continue

            content = item.get("content")

            # ------------------------------------------------
            # content = 문자열
            # ------------------------------------------------

            if isinstance(content, str):

                results.append({
                    "type": "prompt",
                    "content": content,
                })

            # ------------------------------------------------
            # content = 배열
            # ------------------------------------------------

            elif isinstance(content, list):

                for content_item in content:

                    if not isinstance(
                        content_item,
                        dict,
                    ):
                        continue

                    content_type = content_item.get(
                        "type"
                    )

                    # ----------------------------------------
                    # Prompt
                    # ----------------------------------------

                    if content_type == "input_text":

                        results.append({
                            "type": "prompt",
                            "content": content_item.get(
                                "text",
                                "",
                            ),
                        })

                    # ----------------------------------------
                    # File
                    # ----------------------------------------

                    elif content_type == "input_file":

                        results.append({
                            "type": "file",
                            "file_id": content_item.get(
                                "file_id"
                            ),
                            "filename": content_item.get(
                                "filename"
                            ),
                        })

                    # ----------------------------------------
                    # Image
                    # ----------------------------------------

                    elif content_type in (
                        "input_image",
                        "image_url",
                    ):

                        results.append({
                            "type": "image",
                            "content": content_item,
                        })

    return results


# ============================================================
# Responses API Gateway Event
# ============================================================

def create_responses_event(
    flow: http.HTTPFlow,
    body: dict,
) -> dict:

    contents = extract_content(body)

    event = {

        "request_id": str(
            uuid.uuid4()
        ),

        "timestamp": (
            datetime.now(
                timezone.utc
            ).isoformat()
        ),

        "client": {
            "ip": get_client_ip(flow),
        },

        "service": "OpenAI",

        "request": {
            "method": flow.request.method,
            "endpoint": flow.request.path,
            "model": body.get("model"),
        },

        "contents": contents,

        "metadata": {
            "content_count": len(contents),
            "body_size": len(
                flow.request.raw_content or b""
            ),
        },
    }

    return event


# ============================================================
# Multipart File Upload 파싱
# ============================================================

def extract_file_upload_info(
    flow: http.HTTPFlow,
) -> dict:

    req = flow.request

    content_type = req.headers.get(
        "content-type",
        "",
    )

    body = req.raw_content or b""

    filename = None
    purpose = None
    file_bytes = None

    # --------------------------------------------------------
    # Content-Type 확인
    # --------------------------------------------------------

    if not content_type.startswith(
        "multipart/form-data"
    ):

        return {
            "filename": None,
            "purpose": None,
            "body_size": len(body),
            "file_bytes": None,
        }

    # --------------------------------------------------------
    # Multipart parser용 가상 MIME Header 생성
    # --------------------------------------------------------

    try:

        raw_message = (
            b"Content-Type: "
            + content_type.encode()
            + b"\r\n"
            + b"\r\n"
            + body
        )

        message = BytesParser(
            policy=default
        ).parsebytes(
            raw_message
        )

        # ----------------------------------------------------
        # Multipart 각 Part 처리
        # ----------------------------------------------------

        for part in message.iter_parts():

            disposition = part.get(
                "Content-Disposition",
                "",
            )

            part_name = part.get_param(
                "name",
                header="Content-Disposition",
            )

            # -----------------------------------------------
            # purpose
            # -----------------------------------------------

            if part_name == "purpose":

                payload = part.get_payload(
                    decode=True
                )

                if payload:

                    purpose = payload.decode(
                        "utf-8",
                        errors="ignore",
                    ).strip()

            # -----------------------------------------------
            # file
            # -----------------------------------------------

            elif part_name == "file":

                filename = part.get_filename()

                file_bytes = part.get_payload(
                    decode=True
                )

    except Exception as e:

        print(
            "❌ Multipart parsing failed:",
            e,
        )

    return {
        "filename": filename,
        "purpose": purpose,
        "body_size": len(body),
        "file_bytes": file_bytes,
    }


# ============================================================
# File Upload Gateway Event
# ============================================================

def create_file_event(
    flow: http.HTTPFlow,
) -> dict:

    file_info = extract_file_upload_info(
        flow
    )

    filename = file_info["filename"]
    file_bytes = file_info["file_bytes"]

    # --------------------------------------------------------
    # 실제 파일 텍스트 추출
    # --------------------------------------------------------

    extracted_text = None

    if file_bytes:

        extracted_text = extract_file_text(
            filename,
            file_bytes,
        )

    # --------------------------------------------------------
    # Content 생성
    # --------------------------------------------------------

    file_content = {
        "type": "file",
        "filename": filename,
        "purpose": file_info["purpose"],
    }

    # 텍스트 추출 성공 시 content 추가
    if extracted_text is not None:

        file_content["content"] = extracted_text

    # --------------------------------------------------------
    # Event 생성
    # --------------------------------------------------------

    event = {

        "request_id": str(
            uuid.uuid4()
        ),

        "timestamp": (
            datetime.now(
                timezone.utc
            ).isoformat()
        ),

        "client": {
            "ip": get_client_ip(flow),
        },

        "service": "OpenAI",

        "request": {
            "method": flow.request.method,
            "endpoint": flow.request.path,
        },

        "contents": [
            file_content
        ],

        "metadata": {
            "content_count": 1,
            "body_size": file_info["body_size"],
        },
    }

    return event


# ============================================================
# 다음 모듈로 Event 전달
# ============================================================

def send_to_next_module(
    event: dict,
):

    try:

        session = requests.Session()

        # 시스템 Proxy 환경변수 무시
        # Gateway → localhost 통신이
        # 다시 Gateway를 거치지 않도록 함
        session.trust_env = False

        response = session.post(
            NEXT_MODULE_URL,
            json=event,
            timeout=3,
        )

        print(
            "\n===== NEXT MODULE RESPONSE ====="
        )

        print(
            "Status:",
            response.status_code,
        )

        print(
            "Body:",
            response.text,
        )

    except requests.RequestException as e:

        print(
            "\n❌ Failed to send Gateway Event"
        )

        print(
            "Error:",
            e,
        )


# ============================================================
# Event 출력
# ============================================================

def print_gateway_event(
    event: dict,
):

    print(
        "\n===== GATEWAY EVENT ====="
    )

    print(
        json.dumps(
            event,
            ensure_ascii=False,
            indent=2,
        )
    )


# ============================================================
# mitmproxy Request Hook
# ============================================================

def request(
    flow: http.HTTPFlow,
):

    req = flow.request

    content_type = req.headers.get(
        "content-type",
        "",
    )

    # ========================================================
    # 1. OpenAI File Upload
    # ========================================================

    if (
        req.host == OPENAI_HOST
        and req.method == "POST"
        and req.path == OPENAI_FILES_ENDPOINT
        and content_type.startswith(
            "multipart/form-data"
        )
    ):

        print(
            "\n===== OPENAI FILE UPLOAD DETECTED ====="
        )

        print(
            "Host:",
            req.host,
        )

        print(
            "Method:",
            req.method,
        )

        print(
            "Path:",
            req.path,
        )

        print(
            "Content-Type:",
            content_type,
        )

        # ----------------------------------------------------
        # File Event 생성
        # ----------------------------------------------------

        event = create_file_event(
            flow
        )

        print_gateway_event(
            event
        )

        # ----------------------------------------------------
        # 다음 모듈 전달
        # ----------------------------------------------------

        send_to_next_module(
            event
        )

        return

    # ========================================================
    # 2. OpenAI Responses API
    # ========================================================

    if not (
        req.host == OPENAI_HOST
        and req.method == "POST"
        and req.path == OPENAI_RESPONSES_ENDPOINT
        and content_type.startswith(
            "application/json"
        )
    ):

        return

    print(
        "\n===== OPENAI REQUEST DETECTED ====="
    )

    print(
        "Host:",
        req.host,
    )

    print(
        "Method:",
        req.method,
    )

    print(
        "Path:",
        req.path,
    )

    print(
        "Content-Type:",
        content_type,
    )

    # --------------------------------------------------------
    # JSON Body 파싱
    # --------------------------------------------------------

    try:

        body = json.loads(
            req.get_text()
        )

    except (
        json.JSONDecodeError,
        TypeError,
    ):

        print(
            "❌ JSON parsing failed"
        )

        return

    # --------------------------------------------------------
    # Gateway Event 생성
    # --------------------------------------------------------

    event = create_responses_event(
        flow,
        body,
    )

    print_gateway_event(
        event
    )

    # --------------------------------------------------------
    # 다음 모듈 전달
    # --------------------------------------------------------

    send_to_next_module(
        event
    )